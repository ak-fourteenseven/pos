"""
商品與分類 API 路由
"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import Optional, List
from io import BytesIO
import openpyxl
from backend.database.config import get_db, generate_id
from backend.models.product import Product, Category
from backend.models.order import Order, OrderItem, OrderStatus

router = APIRouter()


class ProductCreate(BaseModel):
    name: str
    category_id: str
    price: float
    cost: float = 0
    tax_rate: float = 0
    description: str = ""


class ProductUpdate(BaseModel):
    name: Optional[str] = None
    category_id: Optional[str] = None
    price: Optional[float] = None
    cost: Optional[float] = None
    tax_rate: Optional[float] = None
    is_available: Optional[bool] = None
    description: Optional[str] = None


@router.get("/categories")
def list_categories(db: Session = Depends(get_db)):
    categories = db.query(Category).order_by(Category.sort_order).all()
    return {"categories": [c.to_dict() for c in categories]}


@router.get("")
def list_products(category_id: Optional[str] = None, db: Session = Depends(get_db)):
    q = db.query(Product).order_by(Product.sort_order)
    if category_id:
        q = q.filter(Product.category_id == category_id)
    products = q.all()
    return {"products": [p.to_dict() for p in products]}


@router.post("")
def create_product(data: ProductCreate, db: Session = Depends(get_db)):
    product = Product(
        name=data.name,
        category_id=data.category_id,
        price=data.price,
        cost=data.cost,
        tax_rate=data.tax_rate,
        description=data.description,
    )
    db.add(product)
    db.commit()
    db.refresh(product)
    return {"product": product.to_dict()}


@router.put("/{product_id}")
def update_product(product_id: str, data: ProductUpdate, db: Session = Depends(get_db)):
    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="商品不存在")
    if data.name is not None:
        product.name = data.name
    if data.category_id is not None:
        product.category_id = data.category_id
    if data.price is not None:
        product.price = data.price
    if data.cost is not None:
        product.cost = data.cost
    if data.tax_rate is not None:
        product.tax_rate = data.tax_rate
    if data.is_available is not None:
        product.is_available = data.is_available
    if data.description is not None:
        product.description = data.description
    db.commit()
    db.refresh(product)
    return {"product": product.to_dict()}


@router.delete("/{product_id}")
def delete_product(product_id: str, db: Session = Depends(get_db)):
    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="商品不存在")

    # 檢查是否有進行中的訂單包含此商品
    active_orders = (
        db.query(Order)
        .join(OrderItem, Order.id == OrderItem.order_id)
        .filter(
            OrderItem.product_id == product_id,
            Order.status == OrderStatus.OPEN,
        )
        .all()
    )
    if active_orders:
        order_numbers = "、".join([o.order_number for o in active_orders])
        raise HTTPException(
            status_code=409,
            detail=f"商品「{product.name}」存在進行中的訂單（{order_numbers}），無法刪除",
        )

    db.delete(product)
    db.commit()
    return {"message": f"商品「{product.name}」已刪除"}


# ===== Excel 匯入 =====

EXCEL_HEADERS = ["名稱", "分類", "售價", "成本", "描述"]


@router.post("/import")
async def import_products(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="僅支援 .xlsx / .xls 檔案")

    # 載入分類對照表
    categories = db.query(Category).all()
    cat_name_to_id = {c.name: c.id for c in categories}
    existing_names = {p.name for p in db.query(Product.name).all()}

    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(BytesIO(contents))
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="無法讀取 Excel 檔案")

    rows = list(ws.iter_rows(min_row=2, values_only=True))  # 跳過表頭
    success = 0
    failures: List[dict] = []

    for idx, row in enumerate(rows, start=2):
        if not row or all(c is None for c in row):
            continue

        name = str(row[0]).strip() if row[0] else ""
        cat_name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        price_raw = row[2] if len(row) > 2 else None
        cost_raw = row[3] if len(row) > 3 else None
        desc = str(row[4]).strip() if len(row) > 4 and row[4] else ""

        # 校驗
        row_errors = []
        if not name:
            row_errors.append("名稱為空")
        elif name in existing_names:
            row_errors.append(f"名稱「{name}」已存在")
        else:
            existing_names.add(name)  # 同一批次內也防止重複

        if not cat_name:
            row_errors.append("分類為空")
        elif cat_name not in cat_name_to_id:
            row_errors.append(f"分類「{cat_name}」不存在")

        try:
            price = float(price_raw) if price_raw is not None else 0
        except (ValueError, TypeError):
            price = None
            row_errors.append("售價格式錯誤")

        try:
            cost = float(cost_raw) if cost_raw is not None else 0
        except (ValueError, TypeError):
            cost = 0

        if price is None or price < 0:
            row_errors.append("售價無效")

        if row_errors:
            failures.append({"row": idx, "name": name, "errors": row_errors})
            continue

        # 寫入
        product = Product(
            id=generate_id("P"),
            name=name,
            category_id=cat_name_to_id[cat_name],
            price=price,
            cost=cost,
            tax_rate=0,
            description=desc,
        )
        db.add(product)
        success += 1

    db.commit()
    return {
        "total": success + len(failures),
        "success": success,
        "failed": len(failures),
        "failures": failures,
    }


# ===== Excel 匯出 =====

@router.get("/export")
def export_products(db: Session = Depends(get_db)):
    products = db.query(Product).order_by(Product.category_id, Product.sort_order).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "商品列表"

    headers = ["商品編號", "名稱", "分類", "售價", "成本", "狀態", "描述"]
    ws.append(headers)

    for p in products:
        cat_name = p.category.name if p.category else ""
        ws.append([
            p.id,
            p.name,
            cat_name,
            p.price,
            p.cost,
            "上架" if p.is_available else "下架",
            p.description,
        ])

    # 調整欄寬
    col_widths = [14, 22, 10, 10, 10, 8, 30]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=products.xlsx"},
    )


# ===== 匯入模板下載 =====

@router.get("/template")
def download_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "商品匯入模板"

    ws.append(EXCEL_HEADERS)

    # 範例數據
    examples = [
        ["凍檸茶", "飲品", 22, 5, "港式經典凍檸茶"],
        ["叉燒飯", "燒味", 48, 18, "招牌蜜汁叉燒"],
    ]
    for row in examples:
        ws.append(row)

    # 調整欄寬
    col_widths = [16, 12, 10, 10, 30]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # 給「分類」欄添加資料驗證提示
    ws["B2"].comment = openpyxl.comments.Comment("請填寫系統中已存在的分類名稱", "系統")
    ws["B2"].comment.width = 280
    ws["B2"].comment.height = 40

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=product_import_template.xlsx"},
    )
